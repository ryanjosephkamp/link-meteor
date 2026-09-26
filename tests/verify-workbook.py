#!/usr/bin/env python3
"""Independently inspect the generated OOXML workbook and, when present, open it with openpyxl.

A workbook whose expected JSON has "formatted": true is the 0.3.0 Export panel workbook. For it,
openpyxl also checks the bold frozen header, the filter range, fitted column widths, hyperlinks
whose targets are the exact cell text, text-only cells and the About sheet. openpyxl is a library
reader, not Microsoft Excel.
"""
import json
import math
import re
import sys
import unicodedata
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
LINKABLE = re.compile(r"^(https?://|mailto:)\S+$", re.IGNORECASE)
INFORMATIONAL_ABOUT_LABELS = {"Cells", "Clickable links"}


def fail(message):
    raise SystemExit(message)


def decode_ooxml(value):
    # One pass: _x005F_x0001_ becomes the literal text _x0001_, not a control.
    return re.sub(r"_x([0-9A-Fa-f]{4})_", lambda match: chr(int(match.group(1), 16)), value)


def verify(xlsx_path, expected_path):
    expected = json.loads(expected_path.read_text(encoding="utf-8"))
    with zipfile.ZipFile(xlsx_path) as archive:
        bad = archive.testzip()
        if bad:
            fail(f"ZIP CRC failed: {bad}")
        required = {
            "[Content_Types].xml", "_rels/.rels", "xl/workbook.xml",
            "xl/_rels/workbook.xml.rels", "xl/worksheets/sheet1.xml",
        }
        if expected.get("formatted"):
            required |= {"xl/styles.xml", "xl/worksheets/sheet2.xml"}
        if not required.issubset(archive.namelist()):
            fail("Missing OOXML workbook entries")
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                ElementTree.fromstring(archive.read(name))
                if re.search(rb"<f[ >]", archive.read(name)):
                    fail(f"Formula element in {name}")
        sheet = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        other_sheets = [ElementTree.fromstring(archive.read(name)) for name in archive.namelist()
                        if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name) and name != "xl/worksheets/sheet1.xml"]
    for extra in other_sheets:
        for cell in extra.iter(f"{NS}c"):
            if cell.get("t") != "inlineStr":
                fail(f"Unsafe/non-string cell at {cell.get('r')} on a later sheet")
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    structural_rows = []
    for row in sheet.findall(f"./{namespace}sheetData/{namespace}row"):
        cells = []
        for cell in row.findall(f"./{namespace}c"):
            if cell.get("t") != "inlineStr":
                fail(f"Unsafe/non-string cell at {cell.get('r')}")
            value = cell.find(f"./{namespace}is/{namespace}t")
            cells.append(decode_ooxml(value.text or "") if value is not None else "")
        structural_rows.append(cells)
    if structural_rows != expected["rows"]:
        fail(f"OOXML cell mismatch: {structural_rows!r} != {expected['rows']!r}")

    result = {"zip_crc": "pass", "ooxml_structure": "pass", "rows": len(structural_rows) - 1}
    try:
        from openpyxl import load_workbook
    except ImportError:
        result["openpyxl"] = "unavailable; structural ZIP/XML verification only"
    else:
        book = load_workbook(xlsx_path, read_only=True, data_only=False)
        try:
            raw_rows = [[str(cell) if cell is not None else "" for cell in row] for row in book.active.values]
        finally:
            book.close()
        if raw_rows == expected["rows"]:
            result["openpyxl"] = "opened; exact cell values matched"
        elif [[decode_ooxml(cell) for cell in row] for row in raw_rows] == expected["rows"]:
            result["openpyxl"] = "opened; OOXML escaped strings required decoding for exact match"
        else:
            fail(f"openpyxl cell mismatch: {raw_rows!r} != {expected['rows']!r}")
        if expected.get("formatted"):
            result.update(check_formatted(xlsx_path, expected["rows"], expected["urlColumns"], expected.get("about")))
    print(json.dumps(result, sort_keys=True))


def display_width(value):
    lines = str(value).replace("\r\n", "\n").replace("\r", "\n").split("\n")
    return max(sum(2 if unicodedata.east_asian_width(char) in ("W", "F") else 1 for char in line) for line in lines)


def local_and_utc(seconds):
    local = datetime.fromtimestamp(seconds).astimezone()
    offset = local.strftime("%z")
    return (local.strftime("%Y-%m-%d %H:%M:%S") + f" (UTC{offset[:3]}:{offset[3:]})",
            datetime.fromtimestamp(seconds, timezone.utc).strftime("%Y-%m-%d %H:%M:%S") + " UTC")


def check_formatted(xlsx_path, expected_rows, url_columns, about=None):
    """openpyxl checks for the formatted workbook. about: expected About values, or None to check
    only its shape (a downloaded workbook's export time is not known in advance)."""
    from openpyxl import load_workbook
    book = load_workbook(xlsx_path, read_only=False, data_only=False)
    if book.sheetnames != ["Links", "About"]:
        fail(f"Unexpected sheets: {book.sheetnames}")
    if book.active.title != "Links":
        fail("The Links sheet is not the one that opens first")
    links, about_sheet = book["Links"], book["About"]
    values = [["" if cell.value is None else str(cell.value) for cell in row] for row in links.iter_rows()]
    # openpyxl 3.0 leaves OOXML _xHHHH_ escapes encoded; decode them once, as verify() does.
    if values != expected_rows and [[decode_ooxml(cell) for cell in row] for row in values] != expected_rows:
        fail(f"Formatted cell mismatch: {values!r} != {expected_rows!r}")
    header = expected_rows[0]
    last = f"{links.cell(row=1, column=len(header)).column_letter}{len(expected_rows)}"
    checks = {}
    checks["header_bold"] = all(links.cell(row=1, column=index + 1).font.b for index in range(len(header)))
    checks["body_not_bold"] = not any(cell.font.b for row in links.iter_rows(min_row=2) for cell in row)
    checks["freeze_panes"] = links.freeze_panes
    checks["auto_filter"] = links.auto_filter.ref
    widths, fitted = [], True
    for index, label in enumerate(header):
        letter = links.cell(row=1, column=index + 1).column_letter
        width = links.column_dimensions[letter].width
        widest = max([math.ceil(display_width(label) * 1.1)] + [display_width(row[index]) for row in expected_rows[1:]])
        target = min(80, max(8, widest + 2))
        widths.append(width)
        fitted = fitted and width is not None and abs(width - target) <= 1
    checks["column_widths"] = widths
    checks["widths_fitted"] = fitted
    cells = [cell for sheet in (links, about_sheet) for row in sheet.iter_rows() for cell in row]
    checks["all_cells_strings"] = all(cell.data_type == "s" for cell in cells)
    checks["all_cells_text_format"] = all(cell.number_format == "@" for cell in cells)
    linked, exact, missing, misplaced = 0, True, [], []
    for row in links.iter_rows(min_row=2):
        for cell in row:
            column = header[cell.column - 1]
            text = decode_ooxml("" if cell.value is None else str(cell.value))
            if cell.hyperlink is not None:
                linked += 1
                exact = exact and cell.hyperlink.target == text
                if column not in url_columns:
                    misplaced.append(cell.coordinate)
            elif column in url_columns and LINKABLE.match(text) and len(text) <= 2079:
                missing.append(cell.coordinate)
    checks["hyperlinks"] = linked
    checks["hyperlink_targets_equal_cell_text"] = exact
    checks["unlinked_web_or_mail_cells"] = missing
    checks["links_outside_url_columns"] = misplaced
    about_rows = [["" if cell.value is None else str(cell.value) for cell in row] for row in about_sheet.iter_rows()]
    checks["about_labels_bold"] = all(about_sheet.cell(row=index + 1, column=1).font.b for index in range(len(about_rows)))
    core = [row for row in about_rows if row[0] not in INFORMATIONAL_ABOUT_LABELS]
    labels = [row[0] for row in core]
    checks["about_rows"] = about_rows
    shape = (labels[:5] == ["Exported (local time)", "Exported (UTC)", "Collection", "Links", "View"]
             and labels[-2:] == ["Columns", "Link Meteor version"]
             and (labels[5:-2] == ["Filters"] and core[5][1] == "None" or set(labels[5:-2]) == {"Filter"}))
    if about is not None:
        local, utc = local_and_utc(about["exportedAtUtcSeconds"])
        filters = [["Filter", item] for item in about["filters"]] or [["Filters", "None"]]
        wanted = [["Exported (local time)", local], ["Exported (UTC)", utc], ["Collection", about["collection"]],
                  ["Links", str(about["count"])], ["View", about["view"]], *filters,
                  ["Columns", about["columns"]], ["Link Meteor version", about["version"]]]
        checks["about_matches"] = core == wanted
    else:
        checks["about_matches"] = shape
    book.close()
    ok = (checks["header_bold"] and checks["body_not_bold"] and checks["freeze_panes"] == "A2"
          and checks["auto_filter"] == f"A1:{last}" and checks["widths_fitted"] and checks["all_cells_strings"]
          and checks["all_cells_text_format"] and checks["hyperlink_targets_equal_cell_text"]
          and not missing and not misplaced and checks["about_labels_bold"] and shape and checks["about_matches"])
    if not ok:
        fail(f"Formatted workbook check failed: {json.dumps(checks, ensure_ascii=False)}")
    return {"formatted": "pass", "formatted_checks": checks}


if __name__ == "__main__":
    if len(sys.argv) != 3:
        fail("usage: verify-workbook.py XLSX EXPECTED_JSON")
    verify(Path(sys.argv[1]), Path(sys.argv[2]))
